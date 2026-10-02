"""El manual de usuario en Word y PDF, a partir de 07_manual_de_usuario.md.

El texto y las capturas se cambian en el .md; este guion solo maqueta:
1. pandoc convierte el .md (con su indice y sus capturas) a Word;
2. python-docx le da formato: A4, estilos, tablas, figuras numeradas, cabecera
   y pie con el numero de pagina;
3. Word actualiza el indice y exporta el PDF (finalizar_word.ps1).

Necesita pandoc (vale el que trae RStudio), python-docx y Word. Desde la raiz:
    python clase/sesion_1_2026-09-27/07_manual/construir_word.py
Las capturas se rehacen con capturar.py.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
CLASE = AQUI.parent
FUENTE = CLASE / "07_manual_de_usuario.md"
SALIDA = CLASE / "07_manual_de_usuario.docx"
PANDOC_RSTUDIO = Path(r"C:\Program Files\RStudio\resources\app\bin\quarto\bin\tools\pandoc.exe")

ANCHO_TEXTO = 16.0   # cm: A4 con margenes de 2,5 cm
# Capturas que no van a todo el ancho (las estrechas y altas, o las de poco texto).
ANCHOS = {"01_acceso.png": 11, "03_barra_lateral.png": 10, "04_hoy.png": 14.5,
          "05_tiempos_duros.png": 14.5, "06_regla_en_el_pasado.png": 13.5, "12_clase01_capital.png": 14}

AZUL, AZUL_MEDIO, GRIS = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x2E, 0x74, 0xB5), RGBColor(0x59, 0x59, 0x59)
LETRA, LETRA_CODIGO = "Calibri", "Consolas"

INDICE = """
```{=openxml}
<w:p><w:r><w:br w:type="page"/></w:r></w:p>
<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr><w:r><w:t>Contenido</w:t></w:r></w:p>
<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r><w:r><w:instrText xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>El indice se completa al abrir el documento en Word.</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
```
"""


def pandoc() -> str:
    ruta = shutil.which("pandoc") or (str(PANDOC_RSTUDIO) if PANDOC_RSTUDIO.exists() else None)
    if not ruta:
        raise SystemExit("Falta pandoc: instale RStudio o pandoc (https://pandoc.org).")
    return ruta


def preparar(md: str) -> tuple[str, str]:
    """(markdown para pandoc, fecha de revision). El titulo y la fecha van a la portada."""
    lineas = md.splitlines()
    assert lineas[0] == "# Manual de usuario", "el manual debe empezar por su titulo"
    fecha = re.search(r"\*(Revisado el [^*]+?)\.?\*", md).group(1)
    cuerpo = [l for l in lineas[1:] if not l.startswith("*Revisado el") and l.strip() != "---"]
    texto = "\n".join(cuerpo)
    texto = texto.replace("# Parte A ·", INDICE + "\n# Parte A ·", 1)

    def ancho(m: re.Match) -> str:
        return f"{m.group(0)}{{width={ANCHOS.get(Path(m.group(2)).name, ANCHO_TEXTO)}cm}}"
    texto = re.sub(r"!\[([^\]]*)\]\(([^)]+\.png)\)", ancho, texto)
    return texto, fecha


def _shd(padre, color: str) -> None:
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), color)
    padre.append(s)


def _letra(estilo, nombre: str | None = None, tam: float | None = None, negrita: bool | None = None,
           color: RGBColor | None = None, cursiva: bool | None = None) -> None:
    f = estilo.font
    if nombre:
        f.name = nombre
        rfonts = estilo.element.get_or_add_rPr().get_or_add_rFonts()
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(a), nombre)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            rfonts.attrib.pop(qn(a), None)
    if tam:
        f.size = Pt(tam)
    if negrita is not None:
        f.bold = negrita
    if cursiva is not None:
        f.italic = cursiva
    if color is not None:
        f.color.rgb = color


def estilos(doc: Document) -> None:
    st = {s.name: s for s in doc.styles if s.name}
    for nombre in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if nombre in st:
            _letra(st[nombre], LETRA, 10.5)
            pf = st[nombre].paragraph_format
            pf.space_after = Pt(2 if nombre == "Compact" else 6)
            pf.line_spacing = 1.12
    _letra(st["Title"], LETRA, 30, True, AZUL)
    st["Title"].paragraph_format.space_before = Pt(120)
    st["Title"].paragraph_format.space_after = Pt(6)
    _letra(st["Subtitle"], LETRA, 15, False, AZUL_MEDIO)
    _letra(st["Date"], LETRA, 11, False, GRIS)
    st["Date"].paragraph_format.space_after = Pt(36)
    for nivel, (tam, color, antes) in {1: (20, AZUL, 0), 2: (14, AZUL_MEDIO, 16), 3: (12, AZUL_MEDIO, 12)}.items():
        h = st[f"Heading {nivel}"]
        _letra(h, LETRA, tam, True, color, False)
        h.paragraph_format.space_before = Pt(antes)
        h.paragraph_format.space_after = Pt(8 if nivel == 1 else 4)
        h.paragraph_format.keep_with_next = True
    st["Heading 1"].paragraph_format.page_break_before = True   # cada parte, en pagina nueva
    _letra(st["TOC Heading"], LETRA, 20, True, AZUL)
    _letra(st["Image Caption"], LETRA, 9, False, GRIS, True)
    st["Image Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st["Image Caption"].paragraph_format.space_after = Pt(12)
    for nombre in ("Captioned Figure", "Figure"):
        st[nombre].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st[nombre].paragraph_format.keep_with_next = True
        st[nombre].paragraph_format.space_before = Pt(6)
    if "Source Code" in st:
        codigo = st["Source Code"]
        _letra(codigo, LETRA_CODIGO, 8)
        codigo.paragraph_format.space_before = Pt(4)
        codigo.paragraph_format.space_after = Pt(8)
        codigo.paragraph_format.line_spacing = 1.0
        _shd(codigo.element.get_or_add_pPr(), "F2F4F7")
    _letra(st["Verbatim Char"], LETRA_CODIGO, 9.5, color=RGBColor(0x24, 0x29, 0x2F))
    _letra(st["Hyperlink"], color=AZUL_MEDIO)
    bloque = st["Block Text"]
    _letra(bloque, LETRA, 10, False, GRIS, True)
    bloque.paragraph_format.left_indent = Cm(0.6)
    borde = OxmlElement("w:pBdr")
    izq = OxmlElement("w:left")
    for k, v in {"w:val": "single", "w:sz": "18", "w:space": "8", "w:color": "2E74B5"}.items():
        izq.set(qn(k), v)
    borde.append(izq)
    bloque.element.get_or_add_pPr().append(borde)


def codigo(doc: Document) -> None:
    """Las letras de los bloques de codigo heredan el estilo del codigo en linea (9,5 pt): a 8 pt
    caben las ordenes largas sin partirse."""
    for p in doc.paragraphs:
        if p.style.name == "Source Code":
            for r in p.runs:
                r.font.size = Pt(8)


def _texto_celda(celda) -> str:
    return " ".join(p.text for p in celda.paragraphs).strip()


def tablas(doc: Document) -> None:
    for t in doc.tables:
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        for viejo in tblPr.findall(qn("w:tblBorders")) + tblPr.findall(qn("w:tblLayout")):
            tblPr.remove(viejo)
        bordes = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{lado}")
            for k, v in {"w:val": "single", "w:sz": "4", "w:space": "0", "w:color": "BFC5CD"}.items():
                b.set(qn(k), v)
            bordes.append(b)
        tblPr.append(bordes)
        disposicion = OxmlElement("w:tblLayout")
        disposicion.set(qn("w:type"), "fixed")
        tblPr.append(disposicion)
        # Anchos segun el texto de cada columna: las de ordenes cortas no se comen la pagina.
        n = len(t.columns)
        largos = [max(min(len(_texto_celda(f.cells[i])), 70) for f in t.rows) for i in range(n)]
        pesos = [max(l, 8) ** 0.75 for l in largos]
        anchos = [Cm(ANCHO_TEXTO * p / sum(pesos)) for p in pesos]
        grid = t._tbl.tblGrid
        for i, col in enumerate(grid.findall(qn("w:gridCol"))):
            col.set(qn("w:w"), str(int(anchos[i].twips)))
        for j, fila in enumerate(t.rows):
            for i, celda in enumerate(fila.cells):
                celda.width = anchos[i]
                if j == 0:
                    _shd(celda._tc.get_or_add_tcPr(), "DCE6F2")
                for p in celda.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        r.font.size = Pt(9)
                        if j == 0:
                            r.font.bold = True


def figuras(doc: Document) -> int:
    """Antepone 'Figura n.' a cada pie de imagen."""
    n = 0
    for p in doc.paragraphs:
        if p.style.name == "Image Caption":
            n += 1
            numero = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            rpr.append(OxmlElement("w:b"))
            numero.append(rpr)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = f"Figura {n}. "
            numero.append(t)
            p._p.insert(1 if p._p.pPr is not None else 0, numero)
    return n


def pagina(doc: Document, fecha: str) -> None:
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.2)
        s.different_first_page_header_footer = True     # la portada, sin cabecera ni pie
        cab = s.header.paragraphs[0]
        cab.text = f"Manual de usuario · cryptoquant · {fecha.lower()}"
        cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in cab.runs:
            r.font.size, r.font.color.rgb, r.font.name = Pt(8), GRIS, LETRA
        pie = s.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for tipo, texto in (("begin", None), (None, " PAGE "), ("end", None)):
            r = pie.add_run()
            r.font.size, r.font.color.rgb, r.font.name = Pt(9), GRIS, LETRA
            if tipo:
                f = OxmlElement("w:fldChar")
                f.set(qn("w:fldCharType"), tipo)
                r._r.append(f)
            else:
                ins = OxmlElement("w:instrText")
                ins.set(qn("xml:space"), "preserve")
                ins.text = texto
                r._r.append(ins)


def main() -> int:
    md, fecha = preparar(FUENTE.read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        (tmp / "manual.md").write_text(md, encoding="utf-8")
        borrador = tmp / "borrador.docx"
        subprocess.run([pandoc(), str(tmp / "manual.md"), "-f", "markdown", "-t", "docx",
                        # Con resaltado (en blanco y negro), los bloques de codigo llevan el estilo
                        # "Source Code", al que se da fondo; sin el, salen como texto normal.
                        "--syntax-highlighting=monochrome", f"--resource-path={CLASE}",
                        "-M", "title=Manual de usuario",
                        "-M", "subtitle=La clase, el programa cryptoquant y el piloto de riesgo",
                        "-M", f"date={fecha}", "-M", "lang=es-PE", "-o", str(borrador)], check=True)
        doc = Document(str(borrador))
        estilos(doc)
        codigo(doc)
        tablas(doc)
        n = figuras(doc)
        pagina(doc, fecha)
        doc.core_properties.title = "Manual de usuario"
        doc.core_properties.subject = "La clase, el programa cryptoquant y el piloto de riesgo"
        maqueta = tmp / "maqueta.docx"
        doc.save(str(maqueta))
        r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                            str(AQUI / "finalizar_word.ps1"), "-Entrada", str(maqueta), "-Salida", str(SALIDA)],
                           capture_output=True, text=True)
        if r.returncode != 0:
            raise SystemExit(f"Word no pudo terminar el documento:\n{r.stdout}\n{r.stderr}")
    print(f"{SALIDA.name}: {n} figuras, {len(doc.tables)} tablas. {r.stdout.strip()}")
    print(f"PDF: {SALIDA.with_suffix('.pdf').name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
